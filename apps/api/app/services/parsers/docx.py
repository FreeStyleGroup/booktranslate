"""DOCX.

Порядок абзацев и таблиц берётся обходом тела документа, а не двумя списками
`document.paragraphs` и `document.tables`: те отдают своё содержимое отдельно,
и таблица, стоящая в середине главы, оказалась бы в конце перевода.

Тип блока определяется по стилю абзаца, а не по внешнему виду: жирный текст
в 16 пунктов заголовком не является, а размеченный стилем «Heading 2» —
является, даже если выглядит как обычный. Стиль задаёт автор осознанно.

Объекты: таблица — со своей сеткой и объединёнными ячейками; рисунок —
изображения абзаца; формула Word (OMML), стоящая отдельным абзацем;
листинг — подряд идущие абзацы стиля кода. Подпись — абзац стиля подписи
рядом с объектом.
"""

import zipfile
from collections.abc import Iterator, Mapping
from pathlib import Path
from typing import Any

import docx
from docx.document import Document as WordDocument
from docx.opc.part import Part
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

from app.models.block import BlockKind
from app.models.book_object import ObjectType
from app.services.parsers.archives import ArchiveTooLargeError, check_archive
from app.services.parsers.base import (
    ImageSource,
    Owner,
    ParsedBlock,
    ParsedItem,
    ParsedObject,
    ParsingError,
)
from app.services.parsers.formulas import omml_payload
from app.services.parsers.labels import EQUATION_NUMBER, bind_captions, normalize_number

# Имена стилей заголовков в русской и английской сборках Word. Сравнение по
# началу строки: у уровней имена «Heading 1», «Заголовок 2» и так далее.
_HEADING_PREFIXES = ("heading", "заголовок", "title", "название")
_TITLE_PREFIXES = ("title", "название")
_CAPTION_PREFIXES = ("caption", "название объекта", "подпись")
_LIST_PREFIXES = ("list paragraph", "абзац списка", "list bullet", "list number", "маркированный")
# Стили листингов. «HTML Preformatted» и «Стандартный HTML» — одно и то же
# в разных сборках: так Word называет моноширинный текст из веб-страницы.
_CODE_PREFIXES = ("code", "source code", "html preformatted", "стандартный html", "код", "листинг")

# Изображение в старой разметке Word (VML) — у этого пространства имён нет
# короткого имени в python-docx.
_VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"


class DocxParser:
    def parse(self, path: Path) -> Iterator[ParsedItem]:
        # DOCX — это ZIP, и приходит он от клиента: пятьдесят мегабайт по
        # правилам загрузки разворачиваются в десятки гигабайт. Оглавление
        # проверяется до того, как за файл возьмётся библиотека.
        try:
            with zipfile.ZipFile(path) as archive:
                check_archive(archive)
        except zipfile.BadZipFile as error:
            raise ParsingError("Файл не читается как DOCX: повреждён архив") from error
        except ArchiveTooLargeError as error:
            raise ParsingError(str(error)) from error

        try:
            document = docx.Document(str(path))
        except Exception as error:  # библиотека бросает собственные типы ошибок
            raise ParsingError(f"Файл не читается как DOCX: {error}") from error

        items = _Body(document).items()
        bind_captions(items)

        yield from items


class _Body:
    """Обход тела документа с листингом, который копится до первого не-кода."""

    def __init__(self, document: WordDocument) -> None:
        self._document = document
        self._items: list[ParsedItem] = []
        self._index = 0
        self._code: list[str] = []
        self._code_style = ""

    def items(self) -> list[ParsedItem]:
        for child in self._document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                self._paragraph(Paragraph(child, self._document))
            elif child.tag == qn("w:tbl"):
                self._flush_code()
                self._table(Table(child, self._document))

        self._flush_code()

        return self._items

    def _paragraph(self, paragraph: Paragraph) -> None:
        style = _style_name(paragraph)

        # Листинг — знак в знак: пустая строка внутри него — тоже его
        # строка, а отступ не срезается.
        if style.strip().lower().startswith(_CODE_PREFIXES):
            if not self._code:
                self._code_style = style
            self._code.append(paragraph.text)
            return

        self._flush_code()

        images = _images(paragraph, self._document)
        if images:
            self._items.append(
                ParsedObject(
                    type=ObjectType.FIGURE, location={"block": self._index}, images=tuple(images)
                )
            )

        text = paragraph.text.strip()
        formulas = _formulas(paragraph, text)

        if formulas:
            number = EQUATION_NUMBER.match(text)
            label = normalize_number(number.group("number")) if number is not None else None
            self._items.extend(
                ParsedObject(
                    type=ObjectType.EQUATION,
                    location={"block": self._index},
                    payload=omml_payload(formula),
                    label=label if len(formulas) == 1 else None,
                )
                for formula in formulas
            )
            # Номер формулы справа от неё — часть формулы, а не абзац.
            if not text or number is not None:
                return

        if not text:
            return

        kind = _kind_of(paragraph)
        place: dict[str, Any] = {"block": self._index, "style": style}

        if kind is BlockKind.HEADING:
            place["level"] = _heading_level(style)

        self._items.append(ParsedBlock(text=text, kind=kind, location=place))
        self._index += 1

    def _table(self, table: Table) -> None:
        cells, grid = _grid(table)
        parsed = ParsedObject(type=ObjectType.TABLE, location={"block": self._index}, payload=grid)
        self._items.append(parsed)

        for (row, column), cell in cells:
            text = cell.text.strip()
            if not text:
                continue

            self._items.append(
                ParsedBlock(
                    text=text,
                    kind=BlockKind.TABLE_CELL,
                    location={"block": self._index, "row": row, "column": column},
                    owner=Owner(object=parsed),
                )
            )

        self._index += 1

    def _flush_code(self) -> None:
        text = "\n".join(self._code).strip("\n")
        self._code = []

        if not text.strip():
            return

        listing = ParsedObject(
            type=ObjectType.CODE, location={"block": self._index}, payload={"language": None}
        )
        self._items.append(listing)
        self._items.append(
            ParsedBlock(
                text=text,
                kind=BlockKind.CODE,
                location={"block": self._index, "style": self._code_style},
                owner=Owner(object=listing),
            )
        )
        self._index += 1


def _grid(table: Table) -> tuple[list[tuple[tuple[int, int], _Cell]], dict[str, Any]]:
    """Ячейки таблицы с местами в сетке — каждая один раз.

    `row.cells` у python-docx повторяет объединённую ячейку на каждом месте
    сетки, которое она занимает: по горизонтали — столько раз, сколько
    столбцов она охватывает, по вертикали — в каждой строке продолжения.
    Текст такой ячейки попал бы в разбор несколько раз. Поэтому обход идёт
    по самим `<w:tc>`: продолжение вертикального объединения (`vMerge`)
    удлиняет ячейку выше, а не становится новой.
    """
    cells: list[tuple[tuple[int, int], _Cell]] = []
    meta: list[dict[str, Any]] = []
    # Столбец → ячейка, растянутая вниз до текущей строки.
    open_cells: dict[int, dict[str, Any]] = {}
    width = 0

    for row_no, row in enumerate(table.rows):
        column = row.grid_cols_before
        header = row._tr.trPr is not None and row._tr.trPr.find(qn("w:tblHeader")) is not None

        for tc in row._tr.tc_lst:
            span = max(tc.grid_span, 1)

            if tc.vMerge == "continue" and column in open_cells:
                open_cells[column]["rowspan"] += 1
            else:
                entry = {
                    "row": row_no,
                    "col": column,
                    "rowspan": 1,
                    "colspan": span,
                    "header": header,
                }
                meta.append(entry)
                cells.append(((row_no, column), _Cell(tc, table)))
                if tc.vMerge == "restart":
                    open_cells[column] = entry
                else:
                    open_cells.pop(column, None)

            column += span

        width = max(width, column + row.grid_cols_after)

    return cells, {"rows": len(table.rows), "cols": width, "cells": meta}


def _images(paragraph: Paragraph, document: WordDocument) -> list[ImageSource]:
    """Изображения абзаца: файлы из `word/media`, на которые он ссылается.

    Изображение по внешней ссылке (`r:link`) в файле не лежит и рисунком
    не считается: оригинала у сервиса нет.
    """
    related = document.part.related_parts
    found: list[ImageSource] = []

    for drawing in paragraph._p.iter(qn("w:drawing")):
        blip = drawing.find(".//" + qn("a:blip"))
        target = blip.get(qn("r:embed")) if blip is not None else None
        described = drawing.find(".//" + qn("wp:docPr"))
        alt = described.get("descr") if described is not None else None
        source = _related_image(related, target, alt)
        if source is not None:
            found.append(source)

    for imagedata in paragraph._p.iter(_VML_IMAGEDATA):
        source = _related_image(related, imagedata.get(qn("r:id")), imagedata.get("title"))
        if source is not None:
            found.append(source)

    return found


def _related_image(
    related: Mapping[str, Part], target: str | None, alt: str | None
) -> ImageSource | None:
    part = related.get(target) if target else None

    if part is None or not str(part.content_type).startswith("image/"):
        return None

    name = str(part.partname)

    return ImageSource(member=name.lstrip("/"), src=name, alt=alt or None)


def _formulas(paragraph: Paragraph, text: str) -> list[etree._Element]:
    """Выносные формулы абзаца.

    Выносная — `m:oMathPara`, или `m:oMath`, вокруг которой нет текста
    (разве что её номер). Формула внутри фразы остаётся во фразе: без неё
    «где r — радиус» теряет смысл.
    """
    paragraph_math: list[etree._Element] = list(paragraph._p.iterchildren(qn("m:oMathPara")))
    if paragraph_math:
        return paragraph_math

    inline: list[etree._Element] = list(paragraph._p.iterchildren(qn("m:oMath")))
    if inline and (not text or EQUATION_NUMBER.match(text) is not None):
        return inline

    return []


def _style_name(paragraph: Paragraph) -> str:
    """Имя стиля абзаца.

    Стиля может не быть вовсе: у абзаца, размеченного прямым форматированием,
    ссылка на стиль пустая, и обращение к имени напрямую роняло бы разбор
    всего документа из-за одного такого абзаца.
    """
    style = paragraph.style

    return (style.name or "") if style is not None else ""


def _kind_of(paragraph: Paragraph) -> BlockKind:
    name = _style_name(paragraph).strip().lower()

    # Подпись проверяется раньше заголовка: в русской сборке Word стиль
    # подписи называется «Название объекта», а «название» — это ещё и
    # стиль заголовка книги. В обратном порядке каждая подпись к рисунку
    # становилась заголовком — и отдельным разделом главы.
    if name.startswith(_CAPTION_PREFIXES):
        return BlockKind.CAPTION

    if name.startswith(_HEADING_PREFIXES):
        return BlockKind.HEADING

    if name.startswith(_LIST_PREFIXES) or _is_numbered(paragraph):
        return BlockKind.LIST_ITEM

    return BlockKind.PARAGRAPH


def _heading_level(style: str) -> int:
    """Уровень заголовка по имени стиля: «Heading 2» и «Заголовок 2» — второй.

    Заголовок документа («Title», «Название») стоит над всеми уровнями и
    получает нулевой: он встречается один раз и главой не является. Стиль
    заголовка без номера — первый уровень: так его и показывает Word.
    """
    parts = style.strip().split()

    if parts and parts[-1].isdigit():
        return int(parts[-1])

    if style.strip().lower().startswith(_TITLE_PREFIXES):
        return 0

    return 1


def _is_numbered(paragraph: Paragraph) -> bool:
    """Абзац входит в нумерованный или маркированный список.

    Проверяется наличие `w:numPr` в свойствах абзаца — это и есть признак
    списка в разметке. Стиль при этом может называться как угодно: списки
    часто размечают прямым форматированием, минуя стили.
    """
    # Публичного доступа к свойствам абзаца python-docx не даёт.
    properties = paragraph._element.pPr
    return properties is not None and properties.find(qn("w:numPr")) is not None
