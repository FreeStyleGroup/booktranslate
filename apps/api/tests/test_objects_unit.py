"""Специальные объекты без базы: разбор, номера, упоминания, суммы."""

import io
import uuid
import zipfile
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml

from app.models.block import BlockKind
from app.models.book_object import ObjectType, ReferenceKind
from app.services.object_extraction import describe, in_chapter, media_type, read_pictures
from app.services.object_references import KnownObject, TextBlock, find_references
from app.services.parsers import ImageSource, ParsedBlock, ParsedSource, collect
from app.services.parsers.docx import DocxParser
from app.services.parsers.formulas import mathml_payload
from app.services.parsers.labels import caption_label, mentions
from app.services.parsers.markup import blocks_from_html
from app.services.parsers.plain import MarkdownParser
from tests.factories import png_bytes

CHAPTER = {"chapter": 0, "href": "OEBPS/ch03.xhtml"}


def html(markup: str) -> ParsedSource:
    return collect(blocks_from_html(markup, dict(CHAPTER)))


def texts(source: ParsedSource) -> list[str]:
    return [block.text for block in source.blocks]


# ---------- HTML и EPUB ----------


def test_figure_is_an_object_with_its_caption_and_file() -> None:
    source = html(
        '<p>Before.</p><figure id="f32"><img src="img/loss%20curve.png" alt="Loss"/>'
        "<figcaption>Figure 3.2. Loss curve</figcaption></figure>"
    )

    [placed] = source.objects
    assert placed.object.type is ObjectType.FIGURE
    assert placed.at == 1
    assert source.blocks[placed.caption or 0].text == "Figure 3.2. Loss curve"
    assert source.blocks[placed.caption or 0].kind is BlockKind.CAPTION
    # Путь — от файла главы и без процентного кодирования: так он лежит в архиве.
    [image] = placed.object.images
    assert image.member == "OEBPS/img/loss curve.png"
    assert image.alt == "Loss"
    assert placed.object.anchors == ("OEBPS/ch03.xhtml#f32",)


def test_table_grid_follows_html_spans_and_cells_know_their_place() -> None:
    source = html(
        "<table><caption>Table 1.2: Results</caption>"
        '<thead><tr><th>A</th><th colspan="2">B</th></tr></thead>'
        '<tr><td rowspan="2">x</td><td>1</td><td>2</td></tr>'
        "<tr><td><p>3</p></td><td></td></tr></table>"
    )

    [table] = source.objects
    grid = table.object.payload
    assert (grid["rows"], grid["cols"]) == (3, 3)
    # Ячейка под растянутой вниз «x» встаёт во второй столбец, а не в первый.
    assert {"row": 2, "col": 1, "rowspan": 1, "colspan": 1, "header": False} in grid["cells"]
    assert {"row": 0, "col": 1, "rowspan": 1, "colspan": 2, "header": True} in grid["cells"]

    cells = [source.blocks[index] for index in table.content]
    assert [(block.text, block.location["row"], block.location["column"]) for block in cells] == [
        ("A", 0, 0),
        ("B", 0, 1),
        ("x", 1, 0),
        ("1", 1, 1),
        ("2", 1, 2),
        ("3", 2, 1),
    ]
    assert source.blocks[table.caption or 0].text == "Table 1.2: Results"


def test_absurd_span_does_not_explode_the_grid() -> None:
    source = html('<table><tr><td colspan="100000">x</td></tr></table>')

    assert source.objects[0].object.payload["cols"] == 100


def test_display_formula_is_an_object_and_leaves_the_text() -> None:
    source = html(
        "<p>Energy:</p>"
        '<p><math display="block"><mi>E</mi><mo>=</mo><mi>m</mi>'
        "<msup><mi>c</mi><mn>2</mn></msup></math> (3.4)</p>"
        "<p>where <math><mi>r</mi></math> is the radius.</p>"
    )

    [equation] = source.objects
    assert equation.object.type is ObjectType.EQUATION
    assert equation.object.label == "3.4"
    assert equation.object.payload["notation"] == "mathml"
    # Ни формула, ни её номер абзацем не стали; формула в строке осталась во фразе.
    assert texts(source) == ["Energy:", "where r is the radius."]


def test_formula_preview_keeps_math_and_drops_everything_else() -> None:
    source = html(
        '<p><math display="block"><mi onclick="alert(1)" href="https://x">a</mi>'
        "<script>alert(2)</script><mo>+</mo><mn>1</mn></math></p>"
    )

    preview = source.objects[0].object.payload["preview"]
    assert preview == [
        "math",
        {"display": "block"},
        ["mi", {}, "a"],
        ["mo", {}, "+"],
        ["mn", {}, "1"],
    ]


def test_normalized_formula_checksum_ignores_layout_whitespace() -> None:
    from bs4 import BeautifulSoup

    tight = BeautifulSoup("<math><mi>a</mi><mo>+</mo></math>", "html.parser").math
    loose = BeautifulSoup("<math>\n  <mi> a </mi>\n  <mo>+</mo>\n</math>", "html.parser").math
    assert tight is not None and loose is not None

    assert mathml_payload(tight)["normalized"] == mathml_payload(loose)["normalized"]
    assert mathml_payload(tight)["raw"] != mathml_payload(loose)["raw"]


def test_listing_object_owns_its_verbatim_block() -> None:
    source = html('<pre class="language-python">\ndef f():\n    return 1\n</pre>')

    [listing] = source.objects
    assert listing.object.payload == {"language": "python"}
    assert [source.blocks[index].text for index in listing.content] == ["def f():\n    return 1"]


def test_svg_labels_are_not_paragraph_text_and_markup_stays_original() -> None:
    source = html('<p>Diagram: <svg viewBox="0 0 10 10"><text>pump</text></svg></p>')

    assert texts(source) == ["Diagram:"]
    [figure] = source.objects
    [image] = figure.object.images
    # Регистр атрибутов исходника сохранён — иначе SVG рисуется не в том масштабе.
    assert image.data == b'<svg viewBox="0 0 10 10"><text>pump</text></svg>'


def test_neighbour_paragraph_becomes_caption_only_when_it_reads_as_one() -> None:
    captioned = html('<p><img src="a.png"/></p><p>Figure 3.2. Loss curve</p>')
    sentence = html('<p><img src="a.png"/></p><p>Figure 3.2 shows the loss.</p>')

    assert captioned.objects[0].caption == 0
    assert captioned.blocks[0].kind is BlockKind.CAPTION
    assert sentence.objects[0].caption is None
    assert sentence.blocks[0].kind is BlockKind.PARAGRAPH


def test_caption_word_must_match_object_type() -> None:
    source = html("<table><tr><td>1</td></tr></table><p>Figure 2.1: Not this table</p>")

    assert source.objects[0].caption is None


def test_block_links_resolve_to_object_anchors() -> None:
    source = html('<p>See <a href="#f1">the figure</a> and <a href="ch04.xhtml#t2">table</a>.</p>')

    assert source.blocks[0].location["links"] == ["OEBPS/ch03.xhtml#f1", "OEBPS/ch04.xhtml#t2"]


# ---------- Markdown ----------


def test_markdown_fence_is_a_listing_with_language(tmp_path: Path) -> None:
    path = tmp_path / "book.md"
    path.write_text("Text.\n\n```python\n  indented()\n```\n", encoding="utf-8")

    source = collect(MarkdownParser().parse(path))

    [listing] = source.objects
    assert listing.object.type is ObjectType.CODE
    assert listing.object.payload == {"language": "python"}
    assert source.blocks[listing.content[0]].text == "  indented()"


# ---------- DOCX ----------


def docx_source(document: docx.document.Document, tmp_path: Path) -> ParsedSource:
    path = tmp_path / "book.docx"
    document.save(str(path))

    return collect(DocxParser().parse(path))


def test_docx_merged_cell_is_one_block_not_a_repeat(tmp_path: Path) -> None:
    """python-docx повторяет объединённую ячейку на каждом месте сетки — разбор нет."""
    document = docx.Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Wide"
    table.cell(1, 2).merge(table.cell(2, 2)).text = "Tall"
    table.cell(1, 0).text = "a"

    source = docx_source(document, tmp_path)

    assert texts(source) == ["Wide", "a", "Tall"]
    grid = source.objects[0].object.payload
    assert {"row": 0, "col": 0, "rowspan": 1, "colspan": 2, "header": False} in grid["cells"]
    assert {"row": 1, "col": 2, "rowspan": 2, "colspan": 1, "header": False} in grid["cells"]
    assert [source.blocks[index].location["column"] for index in source.objects[0].content] == [
        0,
        0,
        2,
    ]


def test_docx_caption_above_table_belongs_to_the_table(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Table 2.1. Pumps", style="Caption")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "P-101"

    source = docx_source(document, tmp_path)

    assert source.objects[0].object.type is ObjectType.TABLE
    assert source.blocks[source.objects[0].caption or 0].text == "Table 2.1. Pumps"


def test_docx_picture_and_formula_are_objects(tmp_path: Path) -> None:
    picture = tmp_path / "dot.png"
    picture.write_bytes(png_bytes())

    document = docx.Document()
    document.add_picture(str(picture))
    document.add_paragraph("Figure 1.1. Dot", style="Caption")
    formula = document.add_paragraph()
    formula._p.append(
        parse_xml(
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:oMath><m:r><m:t>E=mc2</m:t></m:r></m:oMath></m:oMathPara>"
        )
    )
    formula.add_run("(1.2)")

    source = docx_source(document, tmp_path)

    figure, equation = source.objects
    assert figure.object.type is ObjectType.FIGURE
    assert figure.object.images[0].member is not None
    assert figure.object.images[0].member.startswith("word/media/")
    assert source.blocks[figure.caption or 0].text == "Figure 1.1. Dot"
    assert equation.object.type is ObjectType.EQUATION
    assert equation.object.label == "1.2"
    assert equation.object.payload["text"] == "E=mc2"
    # Номер формулы абзацем не стал.
    assert texts(source) == ["Figure 1.1. Dot"]


def test_docx_code_paragraphs_are_one_verbatim_listing(tmp_path: Path) -> None:
    document = docx.Document()
    document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Intro.")
    document.add_paragraph("def f():", style="Code")
    document.add_paragraph("", style="Code")
    document.add_paragraph("    return 1", style="Code")
    document.add_paragraph("After.")

    source = docx_source(document, tmp_path)

    assert texts(source) == ["Intro.", "def f():\n\n    return 1", "After."]
    [listing] = source.objects
    assert listing.object.type is ObjectType.CODE
    assert listing.content == (1,)


# ---------- номера и упоминания ----------


def test_caption_label_reads_word_number_and_punctuation() -> None:
    label = caption_label("Рисунок 3.2 — Кривая потерь")

    assert label is not None
    assert (label.type, label.number, label.punctuated) == (ObjectType.FIGURE, "3.2", True)
    assert caption_label("Table 3.1 shows results") is not None
    assert caption_label("Table 3.1 shows results").punctuated is False  # type: ignore[union-attr]
    assert caption_label("Results") is None


def test_mentions_understand_cases_lists_and_parentheses() -> None:
    found = mentions("На рисунке 3.2 и в табл. 1.4 видно, что Figures 2.1 and 2.2 и формула (5.1)…")

    assert [(item.type, item.number) for item in found] == [
        (ObjectType.FIGURE, "3.2"),
        (ObjectType.TABLE, "1.4"),
        (ObjectType.FIGURE, "2.1"),
        (ObjectType.FIGURE, "2.2"),
        (ObjectType.EQUATION, "5.1"),
    ]


def test_references_prefer_same_chapter_and_refuse_to_guess() -> None:
    here, there, elsewhere = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()

    def figure(chapter: uuid.UUID, label: str) -> KnownObject:
        return KnownObject(
            id=uuid.uuid4(),
            chapter_id=chapter,
            type=ObjectType.FIGURE,
            label=label,
            anchors=(),
            caption_block_id=None,
        )

    local, foreign = figure(here, "1.1"), figure(there, "1.1")
    twin_a, twin_b = figure(there, "2.1"), figure(elsewhere, "2.1")
    block = TextBlock(
        id=uuid.uuid4(), chapter_id=here, text="See Figure 1.1 and Figure 2.1.", links=()
    )

    found = find_references([block], [local, foreign, twin_a, twin_b])

    # 1.1 — своя глава; 2.1 — в двух чужих главах, и связь не выдумывается.
    assert [item.object_id for item in found] == [local.id]


def test_link_beats_label_and_caption_does_not_mention_itself() -> None:
    chapter = uuid.uuid4()
    caption = uuid.uuid4()
    target = KnownObject(
        id=uuid.uuid4(),
        chapter_id=chapter,
        type=ObjectType.TABLE,
        label="4.1",
        anchors=("OEBPS/ch04.xhtml#t41",),
        caption_block_id=caption,
    )
    linked = TextBlock(
        id=uuid.uuid4(),
        chapter_id=chapter,
        text="See Table 4.1.",
        links=("OEBPS/ch04.xhtml#t41",),
    )
    own_caption = TextBlock(id=caption, chapter_id=chapter, text="Table 4.1. Pumps", links=())

    [found] = find_references([linked, own_caption], [target])

    assert found.kind is ReferenceKind.LINK
    assert found.phrase == "Table 4.1"


# ---------- главы, изображения и суммы ----------


def test_opening_picture_belongs_to_the_chapter_it_opens() -> None:
    """Заставка в начале файла главы — этой главы, а не конца предыдущей."""
    first = list(blocks_from_html("<h1>One</h1><p>Text.</p>", {"chapter": 0, "href": "a.xhtml"}))
    second = list(
        blocks_from_html('<img src="opener.png"/><h1>Two</h1>', {"chapter": 1, "href": "b.xhtml"})
    )
    source = collect([*first, *second])

    assert in_chapter(source, 0, 2) == []
    assert [placed.object.type for placed in in_chapter(source, 2, 3)] == [ObjectType.FIGURE]


def test_media_type_is_read_from_bytes() -> None:
    assert media_type(png_bytes()) == "image/png"
    assert media_type(b"\xff\xd8\xff\xe0rest") == "image/jpeg"
    assert media_type(b'<?xml version="1.0"?>\n<svg xmlns="x"/>') == "image/svg+xml"
    # Подложенная под .png страница остаётся двоичным файлом, а не HTML.
    assert media_type(b"<html><script>alert(1)</script></html>") == "application/octet-stream"


def test_pictures_are_read_once_from_the_archive_and_missing_ones_are_honest(
    tmp_path: Path,
) -> None:
    path = tmp_path / "book.epub"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("OEBPS/img/a.png", png_bytes())
    path.write_bytes(buffer.getvalue())

    present = ImageSource(member="OEBPS/img/a.png", src="img/a.png")
    absent = ImageSource(member="OEBPS/img/missing.png", src="img/missing.png")
    inline = ImageSource(data=b"<svg/>")

    found = read_pictures(path, [present, absent, inline])

    assert found[present].media_type == "image/png"
    assert absent not in found
    assert found[inline].media_type == "image/svg+xml"


def test_table_checksum_follows_cell_text_not_caption() -> None:
    def table(cell: str, caption: str) -> ParsedSource:
        return html(f"<table><caption>{caption}</caption><tr><td>{cell}</td></tr></table>")

    def checksum(source: ParsedSource) -> str:
        return describe(source.objects[0], source.blocks, {}).checksum

    base = checksum(table("10 bar", "Table 1.1. Pumps"))

    assert checksum(table("10 bar", "Table 1.1. Pumps, revised")) == base
    assert checksum(table("12 bar", "Table 1.1. Pumps")) != base


def test_missing_picture_is_described_without_bytes() -> None:
    source = html('<img src="lost.png"/>')

    described = describe(source.objects[0], source.blocks, {})

    assert described.payload["images"] == [
        {"sha256": None, "media_type": None, "size": None, "src": "lost.png", "alt": None}
    ]


def test_label_comes_from_caption_of_the_same_kind() -> None:
    source = html('<figure><img src="a.png"/><figcaption>Рис. 3.7. Схема</figcaption></figure>')

    assert describe(source.objects[0], source.blocks, {}).label == "3.7"


def test_parsed_block_owner_survives_collect() -> None:
    source = html("<table><tr><td>a</td></tr></table>")

    assert isinstance(source.blocks[0], ParsedBlock)
    assert source.objects[0].content == (0,)
