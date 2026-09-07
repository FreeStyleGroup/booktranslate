"""Разбор файлов словаря.

Без базы: разбор чужого файла — чистая логика, и ошибки в ней тихие.
Перепутанные колонки дают словарь, где перевод стоит на месте термина, и
замечают это не при загрузке, а на переводе.
"""

import io
import zipfile

import pytest

from app.models.memory import GlossaryEntryKind, GlossaryTermStatus
from app.services.dictionaries import DictionaryError, read_dictionary
from app.services.dictionaries.tabular import read_table
from app.services.dictionaries.tbx import read_tbx


def sources(contents: object) -> dict[str, str]:
    return {term.source_term: term.target_term for term in contents.terms}  # type: ignore[attr-defined]


def test_csv_with_header() -> None:
    data = "source,target,note\nvalve,клапан,запорная арматура\npump,насос,\n".encode()

    contents = read_table(data)

    assert sources(contents) == {"valve": "клапан", "pump": "насос"}
    assert contents.terms[0].note == "запорная арматура"


def test_csv_without_header() -> None:
    """Словарь без заголовка — норма; потерянная первая строка обнаруживается не сразу."""
    contents = read_table("valve,клапан\npump,насос\n".encode())

    assert sources(contents) == {"valve": "клапан", "pump": "насос"}


def test_semicolon_is_recognised() -> None:
    """Так выгружает русский Excel, и запятые внутри текста сбивают Sniffer."""
    data = "source;target;note\nvalve;клапан;деталь, важная\npump;насос;\n".encode()

    contents = read_table(data)

    assert sources(contents) == {"valve": "клапан", "pump": "насос"}


def test_cp1251_is_decoded() -> None:
    contents = read_table("valve,клапан\n".encode("cp1251"))

    assert sources(contents) == {"valve": "клапан"}


def test_row_without_translation_is_skipped_with_reason() -> None:
    contents = read_table("source,target\nvalve,клапан\npump,\n".encode())

    assert sources(contents) == {"valve": "клапан"}
    assert contents.skipped and "нет термина или перевода" in contents.skipped[0]


def test_kind_column_is_understood() -> None:
    data = "source,target,note,kind\nПЛК,PLC,,аббревиатура\n".encode()

    contents = read_table(data)

    assert contents.terms[0].kind is GlossaryEntryKind.ABBREVIATION
    # Аббревиатуру при первом употреблении положено раскрывать.
    assert contents.terms[0].expand_on_first_use is True


def test_empty_file_is_refused() -> None:
    with pytest.raises(DictionaryError):
        read_table(b"   \n")


TBX = """<?xml version="1.0" encoding="UTF-8"?>
<martif type="TBX-Basic" xml:lang="en">
  <text><body>
    <termEntry id="c1">
      <descrip type="definition">A device that controls flow.</descrip>
      <langSet xml:lang="en-US"><tig><term>valve</term></tig></langSet>
      <langSet xml:lang="ru-RU"><tig><term>клапан</term></tig></langSet>
    </termEntry>
    <termEntry id="c2">
      <langSet xml:lang="en-US"><tig><term>pump</term></tig></langSet>
      <langSet xml:lang="de-DE"><tig><term>Pumpe</term></tig></langSet>
    </termEntry>
  </body></text>
</martif>
""".encode()


def test_tbx_pairs_languages() -> None:
    contents = read_tbx(TBX, source_language="en", target_language="ru")

    assert sources(contents) == {"valve": "клапан"}
    assert contents.terms[0].note == "A device that controls flow."
    # Вторая запись без русского — пропущена с причиной, а не молча.
    assert contents.skipped


def test_tbx_matches_language_by_prefix() -> None:
    """В базах пишут `en-US`, а просят `en`; точного совпадения не бывает почти никогда."""
    contents = read_tbx(TBX, source_language="en-GB", target_language="ru")

    assert "valve" in sources(contents)


def test_tbx_skips_forbidden_variants() -> None:
    """«Не используйте это слово» — указание, противоположное работе глоссария."""
    data = """<?xml version="1.0"?>
    <martif><text><body><termEntry>
      <langSet xml:lang="en"><tig><term>valve</term></tig></langSet>
      <langSet xml:lang="ru"><tig>
        <term>вентиль</term><termNote type="administrativeStatus">notRecommended</termNote>
      </tig></langSet>
    </termEntry></body></text></martif>
    """.encode()

    contents = read_tbx(data, source_language="en", target_language="ru")

    assert contents.terms == []


def test_broken_xml_is_refused() -> None:
    with pytest.raises(DictionaryError):
        read_tbx(b"<martif><termEntry>", source_language="en", target_language="ru")


def registry_docx(entries: list[tuple[str, str, str | None]]) -> bytes:
    """Реестр в том виде, в каком его ведут: стили Entry, Definition, Meta."""
    import docx

    document = docx.Document()

    # Стилей с такими именами в шаблоне Word нет — заводим их сами, иначе
    # python-docx откажется применять несуществующий стиль.
    from docx.enum.style import WD_STYLE_TYPE

    for name in ("Entry", "Definition", "Meta"):
        document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    for entry, definition, meta in entries:
        document.add_paragraph(entry, style="Entry")
        document.add_paragraph(definition, style="Definition")
        if meta is not None:
            document.add_paragraph(meta, style="Meta")

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def test_registry_reads_abbreviation_with_expansion() -> None:
    data = registry_docx(
        [
            (
                "CH02-ABBR-004  ·  MM",
                "market maker — маркет-мейкер. Первое: «маркет-мейкер "
                "(market maker, MM)»; далее MM.",
                "аббревиатура · требует унификации по книге · 2.1 · первый CH02-B0008; печ. 20",
            )
        ]
    )

    contents = read_dictionary(data, "реестр.docx", source_language="en", target_language="ru")
    term = contents.terms[0]

    assert term.source_term == "MM"
    assert term.target_term == "маркет-мейкер"
    assert term.kind is GlossaryEntryKind.ABBREVIATION
    assert term.status is GlossaryTermStatus.NEEDS_UNIFICATION
    assert term.expand_on_first_use is True
    # По ссылке запись находят в исходном реестре: спор о термине решается
    # возвратом к этому месту.
    assert term.reference is not None
    assert "CH02-ABBR-004" in term.reference
    assert "CH02-B0008" in term.reference


def test_registry_dash_inside_explanation_is_not_expansion() -> None:
    """«позиция — для знаковой величины q» — это пояснение, а не расшифровка."""
    data = registry_docx(
        [
            (
                "CH02-TERM-005  ·  inventory",
                "запас актива. Количество актива у участника. Альтернативы: позиция — "
                "для знаковой величины q.",
                "риск · требует унификации по книге · Вводный текст · первый CH02-B0002",
            )
        ]
    )

    contents = read_dictionary(data, "реестр.docx", source_language="en", target_language="ru")

    assert contents.terms[0].target_term == "запас актива"


def test_registry_keeps_notation_as_is() -> None:
    """Обозначение не переводится: σ остаётся σ, а описание уходит в примечание."""
    data = registry_docx(
        [
            (
                "CH02-ABBR-012  ·  n",
                "number of market makers — число конкурирующих маркет-мейкеров.",
                "математическое обозначение · предварительно рекомендован · 2.1.1",
            )
        ]
    )

    term = read_dictionary(data, "реестр.docx", source_language="en", target_language="ru").terms[0]

    assert term.kind is GlossaryEntryKind.NOTATION
    assert term.target_term == "n"
    assert term.note is not None and "число конкурирующих" in term.note


def test_registry_marks_untranslated_by_status() -> None:
    """«Оставить без перевода» стоит в колонке статуса, но говорит о разряде."""
    data = registry_docx(
        [
            (
                "CH02-ABBR-003  ·  NASDAQ",
                "National Association of Securities Dealers Automated Quotations — NASDAQ.",
                "имя площадки · оставить без перевода · Вводный текст · первый CH02-B0006",
            )
        ]
    )

    term = read_dictionary(data, "реестр.docx", source_language="en", target_language="ru").terms[0]

    assert term.kind is GlossaryEntryKind.DO_NOT_TRANSLATE
    assert term.target_term == "NASDAQ"
    assert term.status is GlossaryTermStatus.CONFIRMED


def test_registry_skips_editorial_notes() -> None:
    """Заметки о смежных понятиях идут теми же стилями, но записями не являются."""
    data = registry_docx(
        [
            (
                "liquidity trader / informed trader",
                "Первый торгует из-за потребности в ликвидности, второй использует "
                "информационное преимущество.",
                None,
            )
        ]
    )

    contents = read_dictionary(data, "реестр.docx", source_language="en", target_language="ru")

    assert contents.terms == []
    assert "не запись реестра" in contents.skipped[0]


def test_unknown_extension_is_refused() -> None:
    with pytest.raises(DictionaryError):
        read_dictionary(b"...", "glossary.pdf", source_language="en", target_language="ru")


def test_zip_that_is_not_docx_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a document")

    with pytest.raises(DictionaryError):
        read_dictionary(
            buffer.getvalue(), "glossary.docx", source_language="en", target_language="ru"
        )
