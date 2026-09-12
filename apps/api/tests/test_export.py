"""Выгрузка переведённого документа.

Главное здесь — не форматы, а два обещания: недопереведённую книгу не отдают
молча, а DOCX возвращается с сохранённой структурой, а не голым текстом.
"""

import io
import zipfile

import docx
from httpx import AsyncClient

from tests.conftest import requires_database
from tests.factories import (
    Account,
    PdfLine,
    PdfPage,
    create_project,
    epub_bytes,
    pdf_bytes,
    register,
)

MANUAL = (
    b"# Safety\n"
    b"\n"
    b"Open the valve before start.\n"
    b"\n"
    b"- Check the pressure gauge\n"
    b"\n"
    b"Warning! Disconnect the power supply.\n"
)


async def upload(
    client: AsyncClient, account: Account, *, data: bytes, name: str, media: str
) -> str:
    project_id = await create_project(client, account)

    uploaded = await client.post(
        f"/projects/{project_id}/documents",
        headers=account.headers,
        files={"file": (name, data, media)},
    )
    assert uploaded.status_code == 201, uploaded.text
    document_id: str = uploaded.json()["id"]

    parsed = await client.post(f"/documents/{document_id}/parse", headers=account.headers)
    assert parsed.status_code == 200, parsed.text

    return document_id


async def prepare(
    client: AsyncClient,
    account: Account,
    *,
    data: bytes = MANUAL,
    name: str = "manual.md",
    media: str = "text/markdown",
) -> str:
    document_id = await upload(client, account, data=data, name=name, media=media)

    translated = await client.post(f"/documents/{document_id}/translate", headers=account.headers)
    assert translated.status_code == 200, translated.text

    return document_id


def translated_docx() -> bytes:
    """Документ Word со стилями, списком и таблицей."""
    document = docx.Document()
    document.add_heading("Safety instructions", level=1)
    document.add_paragraph("Open the valve before start.")
    document.add_paragraph("Check the pressure gauge", style="List Bullet")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


@requires_database
async def test_document_says_which_formats_it_exports_to(db_client: AsyncClient) -> None:
    """Кнопки выгрузки в кабинете рисуются по этому списку, а не по догадке витрины."""
    account = await register(db_client)

    markdown_id = await upload(
        db_client, account, data=MANUAL, name="manual.md", media="text/markdown"
    )
    docx_id = await upload(
        db_client,
        account,
        data=translated_docx(),
        name="manual.docx",
        media="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    markdown = await db_client.get(f"/documents/{markdown_id}", headers=account.headers)
    docx_document = await db_client.get(f"/documents/{docx_id}", headers=account.headers)

    # Markdown обратно в себя не собирается — формата оригинала в списке
    # нет, зато есть новый документ Word.
    assert markdown.json()["export_formats"] == ["docx", "markdown", "text"]
    # Word — собирается по месту, и он первым: ради него и приходят. Второй
    # «docx» здесь значил бы то же самое и только путал бы.
    assert docx_document.json()["export_formats"] == ["source", "markdown", "text"]


@requires_database
async def test_pdf_is_handed_over_as_a_fresh_word_document(db_client: AsyncClient) -> None:
    """PDF обратно в PDF не вписать; заказчик получает документ Word со
    структурой: заголовок, абзацы, список."""
    account = await register(db_client)
    document_id = await prepare(
        db_client,
        account,
        data=pdf_bytes(
            [
                PdfPage(
                    lines=(
                        PdfLine("Safety", size=18),
                        PdfLine(""),
                        PdfLine("Open the valve before start."),
                        PdfLine(""),
                        PdfLine("- Check the pressure gauge"),
                        PdfLine("- Close the valve"),
                    )
                )
            ]
        ),
        name="manual.pdf",
        media="application/pdf",
    )

    response = await db_client.get(
        f"/documents/{document_id}/export",
        headers=account.headers,
        params={"format": "docx"},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")

    document = docx.Document(io.BytesIO(response.content))
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs if p.text.strip()]

    assert paragraphs[0] == ("Heading 1", "[ru] Safety")
    assert paragraphs[1][1] == "[ru] Open the valve before start."
    assert paragraphs[2] == ("List Bullet", "[ru] Check the pressure gauge")


@requires_database
async def test_untranslated_document_is_not_handed_over(db_client: AsyncClient) -> None:
    """Книга, где половина абзацев на английском, отданная молча, — худший исход."""
    account = await register(db_client)
    document_id = await upload(
        db_client, account, data=MANUAL, name="manual.md", media="text/markdown"
    )

    response = await db_client.get(
        f"/documents/{document_id}/export?format=markdown", headers=account.headers
    )

    assert response.status_code == 409
    assert "Не переведено блоков: 4" in response.json()["detail"]


@requires_database
async def test_draft_says_how_much_is_missing(db_client: AsyncClient) -> None:
    account = await register(db_client)
    document_id = await upload(
        db_client, account, data=MANUAL, name="manual.md", media="text/markdown"
    )

    response = await db_client.get(
        f"/documents/{document_id}/export?format=markdown&allow_untranslated=true",
        headers=account.headers,
    )

    assert response.status_code == 200, response.text
    assert response.headers["X-Untranslated-Blocks"] == "4"
    # Непереведённое ушло исходным текстом, а не пустотой.
    assert "Open the valve before start." in response.text


@requires_database
async def test_markdown_keeps_the_structure(db_client: AsyncClient) -> None:
    account = await register(db_client)
    document_id = await prepare(db_client, account)

    response = await db_client.get(
        f"/documents/{document_id}/export?format=markdown", headers=account.headers
    )

    assert response.status_code == 200, response.text
    body = response.text

    assert response.headers["X-Untranslated-Blocks"] == "0"
    # Заголовок остался заголовком того же уровня, пункт списка — пунктом.
    assert body.startswith("# [ru] Safety")
    assert "\n- [ru] Check the pressure gauge" in body


@requires_database
async def test_export_name_does_not_shadow_the_original(db_client: AsyncClient) -> None:
    account = await register(db_client)
    document_id = await prepare(db_client, account)

    response = await db_client.get(
        f"/documents/{document_id}/export?format=text", headers=account.headers
    )

    assert (
        "manual-%D0%BF%D0%B5%D1%80%D0%B5%D0%B2%D0%BE%D0%B4.txt"
        in (response.headers["Content-Disposition"])
    )


@requires_database
async def test_docx_comes_back_as_docx_with_styles(db_client: AsyncClient) -> None:
    """Собирать Word с нуля нельзя — оформление заказчика осталось бы в оригинале."""
    account = await register(db_client)
    document_id = await prepare(
        db_client,
        account,
        data=translated_docx(),
        name="manual.docx",
        media="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    response = await db_client.get(f"/documents/{document_id}/export", headers=account.headers)

    assert response.status_code == 200, response.text

    result = docx.Document(io.BytesIO(response.content))
    texts = [paragraph.text for paragraph in result.paragraphs if paragraph.text.strip()]

    assert texts[0].startswith("[ru] Safety instructions")
    # Стиль заголовка на месте: перевод вписан в исходный файл, а не собран
    # заново из голого текста.
    assert result.paragraphs[0].style is not None
    assert result.paragraphs[0].style.name.startswith("Heading")

    # Таблица тоже переведена и осталась таблицей.
    assert len(result.tables) == 1
    assert result.tables[0].cell(0, 0).text.startswith("[ru] Parameter")


@requires_database
async def test_docx_from_other_sources_is_a_fresh_document(db_client: AsyncClient) -> None:
    """Из Markdown документ Word собирается с нуля — со структурой, а не текстом."""
    account = await register(db_client)
    document_id = await prepare(db_client, account)

    response = await db_client.get(
        f"/documents/{document_id}/export?format=docx", headers=account.headers
    )

    assert response.status_code == 200, response.text

    document = docx.Document(io.BytesIO(response.content))
    styles = [p.style.name for p in document.paragraphs if p.text.strip()]

    assert styles[0] == "Heading 1"
    assert "List Bullet" in styles


@requires_database
async def test_source_format_without_writer_says_so(db_client: AsyncClient) -> None:
    account = await register(db_client)
    document_id = await prepare(db_client, account)

    response = await db_client.get(
        f"/documents/{document_id}/export?format=source", headers=account.headers
    )

    assert response.status_code == 415
    assert "markdown" in response.json()["detail"]


@requires_database
async def test_long_paragraph_comes_back_whole(db_client: AsyncClient) -> None:
    """Абзац резался ради перевода — в файл он обязан вернуться абзацем."""
    account = await register(db_client)

    sentence = "The hydraulic power unit operates at a very high pressure indeed. "
    data = (sentence * 40).strip().encode("utf-8")

    document_id = await prepare(db_client, account, data=data, name="long.txt", media="text/plain")

    count = (
        await db_client.get(f"/documents/{document_id}/segments", headers=account.headers)
    ).json()["total"]
    assert count > 1, "абзац должен был разрезаться, иначе тест ничего не проверяет"

    response = await db_client.get(
        f"/documents/{document_id}/export?format=text", headers=account.headers
    )

    # Один блок — одна строка: части склеились обратно.
    assert len([line for line in response.text.splitlines() if line.strip()]) == 1


@requires_database
async def test_epub_comes_back_as_a_readable_book(db_client: AsyncClient) -> None:
    """Книга — это архив с обложкой, стилями и описью; собрать её заново нельзя."""
    account = await register(db_client)
    document_id = await prepare(
        db_client,
        account,
        data=epub_bytes(),
        name="book.epub",
        media="application/epub+zip",
    )

    response = await db_client.get(f"/documents/{document_id}/export", headers=account.headers)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/epub+zip")

    with zipfile.ZipFile(io.BytesIO(response.content)) as book:
        names = book.namelist()

        # Первая запись и без сжатия — по ней читалки опознают формат.
        assert names[0] == "mimetype"
        assert book.getinfo("mimetype").compress_type == zipfile.ZIP_STORED

        # Опись и главы на месте, ничего не потерялось.
        assert "META-INF/container.xml" in names
        assert "OEBPS/content.opf" in names

        chapter = book.read("OEBPS/part0002.xhtml").decode("utf-8")

    assert "[ru] Первая глава" in chapter
    # Разметка блока сохранилась: заголовок остался заголовком.
    assert "<h1>[ru] Первая глава</h1>" in chapter


@requires_database
async def test_html_keeps_its_own_markup(db_client: AsyncClient) -> None:
    account = await register(db_client)
    page = (
        b"<html><body><h2 class='chapter'>Safety</h2>"
        b"<p>Open the valve before start.</p>"
        b"<script>alert(1)</script></body></html>"
    )

    document_id = await prepare(db_client, account, data=page, name="page.html", media="text/html")

    response = await db_client.get(f"/documents/{document_id}/export", headers=account.headers)

    assert response.status_code == 200, response.text
    body = response.text

    # Класс и уровень заголовка остались; переведён только текст.
    assert 'class="chapter"' in body
    assert "[ru] Safety" in body
    assert "[ru] Open the valve before start." in body
    # Скрипт разбор не читал — значит и портить его при сборке незачем.
    assert "alert(1)" in body


@requires_database
async def test_export_is_scoped_to_organization(db_client: AsyncClient) -> None:
    owner = await register(db_client)
    document_id = await prepare(db_client, owner)

    stranger = await register(db_client, email="other@example.com", organization_name="Чужие")
    response = await db_client.get(
        f"/documents/{document_id}/export?format=text", headers=stranger.headers
    )

    assert response.status_code == 404
